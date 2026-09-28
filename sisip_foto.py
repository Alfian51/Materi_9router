import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

SRC = r'd:\Modul Kelompok 8\Bahan_Modul_Kel_8_Simple (1).docx'
DST = r'd:\Modul Kelompok 8\Laporan_Modul_Kelompok_8_Lengkap.docx'
LAMP = r'd:\Modul Kelompok 8\lampiran modul 9router'

# (filename, anchor_text_substring, caption)
PLAN = [
    ("install node.png", "Verifikasi instalasi berhasil", "Gambar 1 - Verifikasi instalasi Node.js (node -v dan npm -v)"),
    ("install 9router.png", "INSTALL 9ROUTER GLOBALLY", "Gambar 2 - Instalasi 9Router secara global via CMD (npm install -g 9router)"),
    ("port 9router.jpeg", "JALANKAN 9ROUTER & BUKA WEB UI", "Gambar 3 - Menjalankan 9Router dan membuka Web UI (port 20128)"),
    ("login 9router.jpeg", "LOGIN KE DASHBOARD 9ROUTER", "Gambar 4 - Login dashboard 9Router dengan password default 123456"),
    ("sambung opencode.png", "AMBIL API KEY OPENCODE", "Gambar 5 - Pendaftaran / login dan penyambungan akun OpenCode"),
    ("Screenshot 2026-09-28 090422.png", "Copy API Key", "Gambar 6 - Tampilan API Key OpenCode yang berhasil dibuat"),
    ("API config.jpeg", "CONNECT OPENCODE KE 9ROUTER", "Gambar 7 - Konfigurasi API / Connect provider OpenCode di menu Providers 9Router"),
    ("combo 9router.jpeg", "BUAT COMBO DI 9ROUTER", "Gambar 8 - Pembuatan Combo OpenCode-AI di menu Combo & Vision Adapter"),
    ("menambah apikey 9router.jpeg", "BUAT API KEY DI 9ROUTER", "Gambar 9 - Pembuatan API Key 9Router di menu Endpoint & Key untuk Antigravity"),
    ("mitm proxy.jpeg", "SETUP ANTIGRAVITY DI CLI TOOLS", "Gambar 10 - Setup Antigravity di CLI Tools: paste API Key dan pilih model (MITM Proxy)"),
    ("mitm server.jpeg", "Klik tombol \"Start Server\"", "Gambar 11 - MITM Server Running setelah klik Start Server"),
]

d = docx.Document(SRC)

def find_idx(substr):
    s = substr.lower()
    for i, p in enumerate(d.paragraphs):
        if s in p.text.lower():
            return i
    return None

# insert from bottom to top agar index tidak bergeser
inserts = []
for fname, anchor, caption in PLAN:
    idx = find_idx(anchor)
    if idx is None:
        print(f'anchor tidak ketemu: {anchor}')
        idx = len(d.paragraphs) - 1
    inserts.append((idx, fname, caption))

inserts.sort(key=lambda x: x[0], reverse=True)

for idx, fname, caption in inserts:
    fpath = os.path.join(LAMP, fname)
    if not os.path.exists(fpath):
        print(f'file hilang: {fpath}')
        continue
    ref_para = d.paragraphs[idx]
    # gambar
    p_img = docx.oxml.OxmlElement('w:p')
    ref_para._p.addnext(p_img)
    from docx.text.paragraph import Paragraph
    new_p = Paragraph(p_img, d)
    new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new_p.add_run()
    run.add_picture(fpath, width=Inches(5.5))
    # caption
    p_cap = docx.oxml.OxmlElement('w:p')
    p_img.addnext(p_cap)
    cap_p = Paragraph(p_cap, d)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = cap_p.add_run(caption)
    r2.font.size = Pt(9)
    r2.font.italic = True
    r2.font.name = 'Calibri'

d.save(DST)
print(f'saved: {DST}')
