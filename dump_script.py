import docx
d = docx.Document(r'd:\Modul Kelompok 8\Bahan_Modul_Kel_8_Simple (1).docx')
with open(r'd:\Modul Kelompok 8\dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(d.paragraphs):
        txt = p.text.strip()
        if txt:
            f.write(f'P{i:02d} [{p.style.name}]: {txt}\n')
    f.write('\n=== TABLES ===\n')
    for ti, t in enumerate(d.tables):
        f.write(f'Table {ti} {len(t.rows)}x{len(t.columns)}\n')
        for row in t.rows:
            f.write(' | '.join(c.text.strip()[:80] for c in row.cells) + '\n')
print('done')
